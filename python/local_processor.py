from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import uuid
from collections import Counter
from pathlib import Path
from typing import Any


STOP_WORDS = {
    "about", "after", "again", "also", "among", "because", "before", "being",
    "between", "could", "during", "first", "from", "have", "into", "more",
    "most", "other", "over", "same", "should", "some", "such", "than", "that",
    "their", "there", "these", "they", "this", "those", "through", "under",
    "using", "very", "what", "when", "where", "which", "while", "with", "would",
}

SUPPORTED_EXTENSIONS = {
    ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
    ".gif", ".docx", ".pptx", ".xlsx", ".xlsm", ".csv", ".txt", ".md",
    ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac", ".mp4", ".mov",
    ".webm", ".mkv", ".avi",
}


def text_summary(text: str) -> str:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", re.sub(r"\s+", " ", text).strip())
    candidates = [sentence.strip() for sentence in sentences if len(sentence.strip()) >= 24]
    if not candidates:
        return text.strip()[:180] or "No text was extracted from this section."

    frequencies = Counter(
        token.lower()
        for token in re.findall(r"\b[a-zA-Z][a-zA-Z0-9'-]{2,}\b", text)
        if token.lower() not in STOP_WORDS
    )

    def score(sentence: str) -> tuple[float, int]:
        words = [
            word.lower()
            for word in re.findall(r"\b[a-zA-Z][a-zA-Z0-9'-]{2,}\b", sentence)
            if word.lower() not in STOP_WORDS
        ]
        return (
            sum(frequencies[word] for word in words) / max(len(words), 1),
            -abs(len(sentence) - 160),
        )

    selected = max(candidates[:80], key=score)
    return selected[:180].rstrip() + ("..." if len(selected) > 180 else "")


def concept_tags(text: str, limit: int = 5) -> list[str]:
    words = re.findall(r"\b[a-zA-Z][a-zA-Z0-9'-]{3,}\b", text.lower())
    counts = Counter(word for word in words if word not in STOP_WORDS)
    return [word for word, _ in counts.most_common(limit)] or ["study material"]


def topic_title(file_name: str, text: str) -> str:
    concepts = concept_tags(text, 3)
    if concepts == ["study material"]:
        name = Path(file_name).stem.replace("_", " ").replace("-", " ").strip()
        return name.title() or "Study Material"
    return " · ".join(concept.title() for concept in concepts)


def extract_image_text(image: Any, ocr_engine: Any) -> str:
    import numpy as np

    result = ocr_engine(np.asarray(image))
    return "\n".join(line.strip() for line in result.txts if line.strip()) if result else ""


def get_ocr_engine(engine_holder: list[Any], warnings: list[str]) -> Any:
    if engine_holder:
        return engine_holder[0]
    try:
        from rapidocr import RapidOCR

        engine = RapidOCR()
        engine_holder.append(engine)
        return engine
    except (ImportError, OSError, RuntimeError) as error:
        warnings.append(
            "Local OCR is unavailable. Install the Python requirements to read scanned PDFs and images. "
            f"Details: {error}"
        )
        return None


def extract_pdf(path: Path, warnings: list[str], ocr_holder: list[Any]) -> list[dict[str, str]]:
    import pymupdf

    sections: list[dict[str, str]] = []
    with pymupdf.open(path) as document:
        if document.needs_pass:
            raise ValueError("This PDF is password-protected and cannot be read.")
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if len(text) < 30:
                engine = get_ocr_engine(ocr_holder, warnings)
                if engine:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                    from PIL import Image

                    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                    ocr_text = extract_image_text(image, engine)
                    if ocr_text:
                        text = ocr_text
            if text:
                sections.append({"location": f"Page {page_number}", "text": text})
    return sections


def extract_image(path: Path, warnings: list[str], ocr_holder: list[Any]) -> list[dict[str, str]]:
    from PIL import Image, ImageOps

    with Image.open(path) as original:
        image = ImageOps.exif_transpose(original).convert("RGB")
        engine = get_ocr_engine(ocr_holder, warnings)
        text = extract_image_text(image, engine) if engine else ""
        if not text:
            text = f"Image file, {image.width} by {image.height} pixels. No readable text was detected."
        return [{"location": f"Image · {image.width} × {image.height}", "text": text}]


def extract_docx(path: Path) -> list[dict[str, str]]:
    from docx import Document

    document = Document(path)
    sections: list[dict[str, str]] = []
    parts: list[str] = []
    section_number = 1
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading") and parts:
            sections.append({"location": f"Section {section_number}", "text": "\n".join(parts)})
            section_number += 1
            parts = []
        parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    if parts:
        sections.append({"location": f"Section {section_number}", "text": "\n".join(parts)})
    return sections


def extract_pptx(path: Path) -> list[dict[str, str]]:
    from pptx import Presentation

    presentation = Presentation(path)
    sections = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        if parts:
            sections.append({"location": f"Slide {slide_number}", "text": "\n".join(parts)})
    return sections


def extract_spreadsheet(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sections = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sections.append({"location": f"Sheet · {worksheet.title}", "text": "\n".join(rows)})
    finally:
        workbook.close()
    return sections


def extract_csv(path: Path) -> list[dict[str, str]]:
    import csv

    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as source:
                rows = [
                    " | ".join(cell.strip() for cell in row if cell.strip())
                    for row in csv.reader(source)
                ]
            return [{"location": "CSV data", "text": "\n".join(row for row in rows if row)}]
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode this CSV as UTF-8 or Windows-1252 text.")


def read_text(path: Path) -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return [{"location": "Text", "text": path.read_text(encoding=encoding).strip()}]
        except UnicodeDecodeError:
            continue
    raise ValueError("This text file encoding is not supported.")


def ffmpeg_details(path: Path) -> tuple[str, str]:
    import imageio_ffmpeg

    result = subprocess.run(
        [imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-i", str(path)],
        capture_output=True,
        text=True,
        timeout=45,
        check=False,
    )
    return result.stderr, result.stdout


def extract_media_metadata(path: Path) -> list[dict[str, str]]:
    try:
        details, _ = ffmpeg_details(path)
        duration_match = re.search(r"Duration:\s*(\d{2}:\d{2}:\d{2}(?:\.\d+)?)", details)
        stream_lines = [
            line.strip()
            for line in details.splitlines()
            if "Stream #" in line and ("Audio:" in line or "Video:" in line)
        ]
        duration = duration_match.group(1) if duration_match else "Unknown"
        stream_summary = "; ".join(stream_lines) or "No audio or video stream was detected."
        return [{
            "location": "Media metadata",
            "text": f"Duration: {duration}. {stream_summary}",
        }]
    except (ImportError, OSError, subprocess.SubprocessError) as error:
        raise ValueError(f"Could not read local audio/video metadata: {error}") from error


def transcribe_media(path: Path, warnings: list[str]) -> list[dict[str, str]]:
    from faster_whisper import WhisperModel
    import imageio_ffmpeg

    with __import__("tempfile").TemporaryDirectory(prefix="studypilot-audio-") as directory:
        audio_path = Path(directory) / "audio.wav"
        subprocess.run(
            [
                imageio_ffmpeg.get_ffmpeg_exe(),
                "-y",
                "-i",
                str(path),
                "-vn",
                "-ac",
                "1",
                "-ar",
                "16000",
                "-f",
                "wav",
                str(audio_path),
            ],
            capture_output=True,
            text=True,
            timeout=900,
            check=True,
        )
        model = WhisperModel("tiny", device="cpu", compute_type="int8")
        segments, _ = model.transcribe(str(audio_path), vad_filter=True)
        output = []
        for segment in segments:
            transcript = segment.text.strip()
            if transcript:
                start = max(0, int(segment.start))
                hours, remainder = divmod(start, 3600)
                minutes, seconds = divmod(remainder, 60)
                timestamp = f"{hours}:{minutes:02d}:{seconds:02d}" if hours else f"{minutes}:{seconds:02d}"
                output.append({"location": timestamp, "text": transcript})
        if not output:
            warnings.append(f"No speech was detected in {path.name}.")
        return output


def extract_sections(
    path: Path,
    audio_mode: str,
    warnings: list[str],
    ocr_holder: list[Any],
) -> tuple[list[dict[str, str]], str, str]:
    extension = path.suffix.lower()
    if extension == ".pdf":
        return extract_pdf(path, warnings, ocr_holder), "pdf", "Document pages"
    if extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}:
        return extract_image(path, warnings, ocr_holder), "image", "Image content"
    if extension == ".docx":
        return extract_docx(path), "text", "Document sections"
    if extension == ".pptx":
        return extract_pptx(path), "slide", "Presentation slides"
    if extension in {".xlsx", ".xlsm"}:
        return extract_spreadsheet(path), "text", "Spreadsheet data"
    if extension == ".csv":
        return extract_csv(path), "text", "CSV data"
    if extension in {".txt", ".md"}:
        return read_text(path), "text", "Text sections"
    if extension in {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac", ".mp4", ".mov", ".webm", ".mkv", ".avi"}:
        if audio_mode == "whisper":
            try:
                return transcribe_media(path, warnings), "video", "Timestamped local transcript"
            except (ImportError, OSError, RuntimeError, subprocess.SubprocessError) as error:
                raise ValueError(
                    "Local Whisper transcription is unavailable. Install Python requirements and retry. "
                    f"The first transcription also downloads the small local model. Details: {error}"
                ) from error
        return extract_media_metadata(path), "video", "Audio/video metadata"
    raise ValueError(f"Unsupported file type: {extension or 'no extension'}.")


def build_result(files: list[dict[str, str]], audio_mode: str) -> dict[str, Any]:
    topics: list[dict[str, Any]] = []
    units: list[dict[str, Any]] = []
    warnings: list[str] = []
    ocr_holder: list[Any] = []

    for file in files:
        original_name = Path(file["name"]).name
        path = Path(file["path"])
        try:
            sections, source_type, subtopic = extract_sections(path, audio_mode, warnings, ocr_holder)
            sections = [section for section in sections if section.get("text", "").strip()]
            if not sections:
                warnings.append(f"No readable content was extracted from {original_name}.")
                continue

            combined_text = "\n".join(section["text"] for section in sections)
            title = topic_title(original_name, combined_text)
            topic_id = f"local-python-topic-{uuid.uuid4()}"
            topic_units = []
            for index, section in enumerate(sections, start=1):
                excerpt = section["text"].strip()[:1800]
                location = section["location"]
                topic_units.append({
                    "id": f"local-python-unit-{uuid.uuid4()}",
                    "title": f"{title} · {location}" if len(sections) > 1 else title,
                    "summary": text_summary(excerpt),
                    "full_text_excerpt": excerpt,
                    "topic_id": topic_id,
                    "topic_name": title,
                    "subtopic": subtopic,
                    "concept_tags": concept_tags(excerpt),
                    "difficulty": "medium",
                    "source_type": source_type,
                    "source_file": original_name,
                    "source_location": location,
                })

            topics.append({
                "id": topic_id,
                "title": title,
                "description": text_summary(combined_text),
                "subtopics": [subtopic],
                "concepts": concept_tags(combined_text, 12),
                "estimated_study_minutes": max(15, min(120, len(topic_units) * 5)),
                "prerequisite_ids": [],
                "unit_count": len(topic_units),
            })
            units.extend(topic_units)
        except Exception as error:
            warnings.append(f"Could not process {original_name}: {error}")

    return {
        "topics": topics,
        "units": units,
        "totalTopics": len(topics),
        "totalUnits": len(units),
        "multimediaUnits": sum(unit["source_type"] in {"image", "video"} for unit in units),
        "warnings": warnings,
        "processor": "Python local libraries",
    }


def report_status() -> None:
    checks = {
        "PyMuPDF": "pymupdf",
        "Pillow": "PIL",
        "RapidOCR": "rapidocr",
        "python-docx": "docx",
        "python-pptx": "pptx",
        "openpyxl": "openpyxl",
        "imageio-ffmpeg": "imageio_ffmpeg",
        "faster-whisper": "faster_whisper",
    }
    dependencies = {}
    for label, module in checks.items():
        try:
            __import__(module)
            dependencies[label] = True
        except ImportError:
            dependencies[label] = False
    print(json.dumps({"python": sys.version.split()[0], "dependencies": dependencies}))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if args.check:
        report_status()
        return

    request = json.load(sys.stdin)
    audio_mode = request.get("audioMode", "metadata")
    if audio_mode not in {"metadata", "whisper"}:
        raise ValueError("Audio processing mode must be metadata or whisper.")
    print(json.dumps(build_result(request.get("files", []), audio_mode), ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error)}), file=sys.stderr)
        sys.exit(1)
